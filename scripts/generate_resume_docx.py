from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def add_markdown_to_doc(md_path: Path, output_path: Path) -> None:
    doc = Document()

    lines = md_path.read_text(encoding="utf-8").splitlines()
    for idx, raw in enumerate(lines):
        line = raw.rstrip()

        if not line:
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(line[2:].strip())
            p.style = doc.styles["Title"]
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(line[3:].strip())
            p.style = doc.styles["Heading 1"]
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip())
            p.style = doc.styles["Heading 2"]
            continue

        if line.startswith("#### "):
            p = doc.add_paragraph(line[5:].strip())
            p.style = doc.styles["Heading 3"]
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(line[2:].strip(), style="List Bullet")
            continue

        # Make contact line centered for better resume layout
        if idx == 1:
            p = doc.add_paragraph(line.strip())
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            continue

        doc.add_paragraph(line)

    doc.save(str(output_path))


if __name__ == "__main__":
    workspace = Path(__file__).resolve().parents[1]
    source = workspace / "Shirish_Resume_Complete.md"
    output = workspace / "Shirish_Resume_Complete.docx"
    add_markdown_to_doc(source, output)
